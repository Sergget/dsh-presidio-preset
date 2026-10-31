"""自定义规则 (正则模板 / 代码模板) 与银行卡 Luhn 降级测试。"""
import json
import os
import sys
import tempfile
import unittest
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from presidio_pii.chinese_recognizers import BankCardRecognizer
from presidio_pii.custom_rules import (RulesLoadError, load_custom_rules,
                                       load_json_rules)
from presidio_pii.engine import PresidioPIIEngine


def _engine(rules=None, whitelist=None):
    """测试统一用内存映射库, 避免污染 .pii_presidio_store/sessions.db。"""
    return PresidioPIIEngine(db_path=":memory:", whitelist=whitelist,
                             custom_rules=rules if rules is not None else [])


class LuhnFallbackTest(unittest.TestCase):
    """Luhn 失败的卡号应降级保留 (0.45), 不再裸漏。"""

    def setUp(self):
        self.rec = BankCardRecognizer()

    def test_luhn_pass_high_score(self):
        # 6222600260001072444 Luhn 通过
        rs = self.rec.analyze("6222600260001072444")
        self.assertEqual(len(rs), 1)
        self.assertEqual(rs[0].entity_type, "BANK_CARD")
        self.assertEqual(rs[0].score, 0.80)

    def test_luhn_fail_kept_low_score(self):
        # 6222020200112233445 Luhn 失败, 但仍须识别 (降级), 不得漏检
        rs = self.rec.analyze("6222020200112233445")
        self.assertEqual(len(rs), 1)
        self.assertEqual(rs[0].score, 0.45)

    def test_engine_redacts_luhn_fail(self):
        out = _engine().redact("卡号 6222020200112233445 请核对", "s-luhn")
        self.assertIn("⟦PII:BANK_CARD:", out.text)

    def test_18digit_id_still_wins_over_bank_card(self):
        # 18 位且行政区划位有效: ID_NUMBER 优先于低分 BANK_CARD
        out = _engine().redact("110101199003077890", "s-prio")
        self.assertIn("⟦PII:ID_NUMBER:", out.text)


class RegexRuleTest(unittest.TestCase):
    """JSON 正则规则模板。"""

    def _write(self, tmp, data, name="rules.json"):
        p = Path(tmp) / name
        p.write_text(json.dumps(data, ensure_ascii=False), encoding="utf-8")
        return p

    def test_basic_rule(self):
        with tempfile.TemporaryDirectory() as td:
            p = self._write(td, [{
                "name": "employee_id",
                "entity_type": "EMPLOYEE_ID",
                "pattern": r"工号[:：\s]*([A-Z]\d{6})",
            }])
            e = _engine(load_json_rules(p))
            out = e.redact("员工工号：A123456，请核实", "s-rule1")
            self.assertIn("⟦PII:EMPLOYEE_ID:0001⟧", out.text)
            self.assertNotIn("A123456", out.text)

    def test_context_hint_required(self):
        with tempfile.TemporaryDirectory() as td:
            p = self._write(td, [{
                "name": "member",
                "entity_type": "MEMBER_ID",
                "pattern": r"[A-Z]{3}\d{8}",
                "context_hint": ["会员"],
            }])
            e = _engine(load_json_rules(p))
            self.assertEqual(e.analyze("编号 ABC12345678"), [])
            self.assertEqual(e.analyze("会员 ABC12345678")[0].entity_type, "MEMBER_ID")

    def test_negative_hint(self):
        with tempfile.TemporaryDirectory() as td:
            p = self._write(td, [{
                "name": "member",
                "entity_type": "MEMBER_ID",
                "pattern": r"[A-Z]{3}\d{8}",
                "negative_hint": ["示例"],
            }])
            e = _engine(load_json_rules(p))
            self.assertEqual(e.analyze("示例 ABC12345678"), [])

    def test_invalid_rule_skipped_lenient(self):
        with tempfile.TemporaryDirectory() as td:
            p = self._write(td, [
                {"name": "bad", "entity_type": "X", "pattern": "("},  # 非法正则
                {"name": "good", "entity_type": "Y", "pattern": r"\d{4}"},
            ])
            rules = load_json_rules(p)
            self.assertEqual(len(rules), 1)
            self.assertEqual(rules[0].entity_type, "Y")

    def test_invalid_rule_raises_strict(self):
        with tempfile.TemporaryDirectory() as td:
            p = self._write(td, [
                {"name": "bad", "entity_type": "X", "pattern": "("},
            ])
            with self.assertRaises(RulesLoadError):
                load_json_rules(p, strict=True)

    def test_validator_relative_to_rule_file(self):
        """validator 的 '模块路径:函数名' 按规则文件目录解析, 与 cwd 无关。"""
        with tempfile.TemporaryDirectory() as td:
            (Path(td) / "validators.py").write_text(
                "def ok(s):\n    return True\n", encoding="utf-8")
            p = self._write(td, [{
                "name": "v", "entity_type": "V", "pattern": r"V\d+",
                "validator": "validators.py:ok",
            }])
            old_cwd = os.getcwd()
            with tempfile.TemporaryDirectory() as other:
                os.chdir(other)  # cwd 与规则文件不同目录
                try:
                    rules = load_json_rules(p, strict=True)
                finally:
                    os.chdir(old_cwd)
            self.assertEqual(rules[0].entity_type, "V")


class CodeRuleTest(unittest.TestCase):
    """Python 代码规则模板。"""

    CODE = (
        "ENTITY_TYPE = 'ORDER_NO'\n"
        "SCORE = 0.9\n"
        "def analyze(text):\n"
        "    import re\n"
        "    return [(m.start(), m.end()) for m in re.finditer(r'ORD-\\d{10}', text)]\n"
    )

    def test_code_rule(self):
        with tempfile.TemporaryDirectory() as td:
            p = Path(td) / "order_rule.py"
            p.write_text(self.CODE, encoding="utf-8")
            e = _engine(load_custom_rules(str(p)))
            out = e.redact("订单号 ORD-1234567890 已发货", "s-code")
            self.assertIn("⟦PII:ORDER_NO:0001⟧", out.text)

    def test_bad_coords_dropped(self):
        with tempfile.TemporaryDirectory() as td:
            p = Path(td) / "bad_rule.py"
            p.write_text(
                "ENTITY_TYPE = 'X'\n"
                "def analyze(text):\n"
                "    return [(-5, 3), (0, 999)]\n", encoding="utf-8")
            e = _engine(load_custom_rules(str(p)))
            self.assertEqual(e.analyze("任意文本"), [])

    def test_exception_isolated(self):
        """代码规则抛异常 / 返回 None 不得中断整条脱敏流水线。"""
        with tempfile.TemporaryDirectory() as td:
            p = Path(td) / "boom.py"
            p.write_text(
                "ENTITY_TYPE = 'X'\n"
                "def analyze(text):\n"
                "    raise RuntimeError('boom')\n", encoding="utf-8")
            e = _engine(load_custom_rules(str(p)))
            out = e.redact("张伟的手机号 13800138000", "s-boom")
            self.assertIn("⟦PII:PERSON:", out.text)  # 内置识别照常工作
            self.assertIn("⟦PII:PHONE_NUMBER:", out.text)

    def test_missing_analyze_rejected(self):
        with tempfile.TemporaryDirectory() as td:
            p = Path(td) / "empty.py"
            p.write_text("ENTITY_TYPE = 'X'\n", encoding="utf-8")
            self.assertEqual(load_custom_rules(str(p)), [])
            with self.assertRaises(RulesLoadError):
                load_custom_rules(str(p), strict=True)

    def test_env_var_loading(self):
        from presidio_pii import custom_rules as cr
        with tempfile.TemporaryDirectory() as td:
            p = Path(td) / "r.json"
            p.write_text(json.dumps(
                [{"name": "n", "entity_type": "ENV_T", "pattern": r"\d{6}"}]),
                encoding="utf-8")
            old = os.environ.get(cr._CUSTOM_ENV)
            os.environ[cr._CUSTOM_ENV] = str(p)
            try:
                rules = cr.load_custom_rules()  # 不传参, 走环境变量
                self.assertEqual(rules[0].entity_type, "ENV_T")
            finally:
                if old is None:
                    os.environ.pop(cr._CUSTOM_ENV, None)
                else:
                    os.environ[cr._CUSTOM_ENV] = old


class IntegrationTest(unittest.TestCase):
    """白名单联动、CLI --rules、preset.json 配置。"""

    RULES = [{
        "name": "employee_id",
        "entity_type": "EMPLOYEE_ID",
        "pattern": r"工号[:：\s]*([A-Z]\d{6})",
    }]

    def test_whitelist_exempts_custom_entity(self):
        with tempfile.TemporaryDirectory() as td:
            p = Path(td) / "rules.json"
            p.write_text(json.dumps(self.RULES, ensure_ascii=False), encoding="utf-8")
            e = _engine(load_json_rules(p), whitelist=["A123456"])
            out = e.redact("员工工号：A123456，请核实", "s-wl")
            self.assertNotIn("⟦PII:EMPLOYEE_ID:", out.text)
            self.assertIn("A123456", out.text)

    def test_cli_rules_end_to_end(self):
        from presidio_pii import cli
        with tempfile.TemporaryDirectory() as td:
            td = Path(td)
            (td / "rules.json").write_text(
                json.dumps(self.RULES, ensure_ascii=False), encoding="utf-8")
            src = td / "sample.txt"
            src.write_text("员工工号：A123456，请核实\n", encoding="utf-8")
            out_dir = td / "out"
            rc = cli.main(["redact", str(src), "--rules", str(td / "rules.json"),
                           "-d", str(out_dir)])
            self.assertEqual(rc, 0)
            text = (out_dir / "sample.txt").read_text(encoding="utf-8")
            self.assertIn("⟦PII:EMPLOYEE_ID:0001⟧", text)

    def test_cli_missing_rules_fails_loudly(self):
        from presidio_pii import cli
        with tempfile.TemporaryDirectory() as td:
            td = Path(td)
            src = td / "sample.txt"
            src.write_text("员工工号：A123456\n", encoding="utf-8")
            rc = cli.main(["redact", str(src), "--rules", str(td / "nope.json"),
                           "-d", str(td / "out")])
            self.assertEqual(rc, 1)

    def test_custom_rule_docx_roundtrip(self):
        """自定义规则在 docx handler 上同样生效, 且写回后可无损还原。"""
        from docx import Document
        from presidio_pii import cli
        with tempfile.TemporaryDirectory() as td:
            td = Path(td)
            (td / "rules.json").write_text(
                json.dumps(self.RULES, ensure_ascii=False), encoding="utf-8")
            doc = Document()
            doc.add_paragraph("员工工号：A123456，请核实")
            doc.save(td / "t.docx")
            self.assertEqual(cli.main(["redact", str(td / "t.docx"), "--rules",
                                       str(td / "rules.json"), "-d", str(td / "out")]), 0)
            redacted = Document(td / "out" / "t.docx").paragraphs[0].text
            self.assertIn("⟦PII:EMPLOYEE_ID:0001⟧", redacted)
            self.assertEqual(cli.main(["restore", str(td / "out" / "t.docx"),
                                       "-d", str(td / "restored")]), 0)
            restored = Document(td / "restored" / "t.docx").paragraphs[0].text
            self.assertEqual(restored, "员工工号：A123456，请核实")

    def test_hooks_read_preset_custom_rules(self):
        """preset.json 的 config.custom_rules_path 应生效 (文档承诺的接入面)。"""
        from presidio_pii import hooks
        with tempfile.TemporaryDirectory() as td:
            td = Path(td)
            (td / "rules.json").write_text(
                json.dumps(self.RULES, ensure_ascii=False), encoding="utf-8")
            preset = td / "preset.json"
            preset.write_text(json.dumps({
                "name": "presidio-pii",
                "config": {"custom_rules_path": "rules.json"},
            }, ensure_ascii=False), encoding="utf-8")
            env = {"PII_DSH_PRESET": str(preset), "PII_DSH_DB_PATH": ":memory:",
                   "PII_DSH_LOG_PATH": "0"}
            old = {k: os.environ.get(k) for k in env}
            os.environ.update(env)
            hooks.reset_singleton()
            try:
                out = hooks.before_llm("员工工号：A123456，请核实", "conv-preset")
                self.assertIn("⟦PII:EMPLOYEE_ID:", out)
            finally:
                hooks.reset_singleton()
                for k, v in old.items():
                    if v is None:
                        os.environ.pop(k, None)
                    else:
                        os.environ[k] = v


if __name__ == "__main__":
    unittest.main()
