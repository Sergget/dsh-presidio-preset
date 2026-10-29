"""端到端流程测试: hooks 日志、DSH 模拟器、真实文档检测工具、CLI。"""
import csv
import json
import os
import sys
from pathlib import Path

import pytest

sys.path.insert(0, str(Path(__file__).parent.parent))
sys.path.insert(0, str(Path(__file__).parent.parent / "tools"))

from presidio_pii import hooks


@pytest.fixture()
def hook_env(tmp_path, monkeypatch):
    """将 hooks 的存储/日志重定向到临时目录并重置单例。"""
    monkeypatch.setenv("PII_DSH_DB_PATH", str(tmp_path / "sessions.db"))
    monkeypatch.setenv("PII_DSH_LOG_PATH", str(tmp_path / "hook_log.jsonl"))
    hooks.reset_singleton()
    yield tmp_path
    hooks.reset_singleton()


# --------------------------------------------------------------------------- #
# hooks + journal
# --------------------------------------------------------------------------- #
def test_hooks_before_after_logging(hook_env):
    text = "张伟的身份证是110101199003077890"
    red = hooks.before_llm(text, conversation_id="conv-1")
    assert "张伟" not in red and "110101199003077890" not in red
    assert "⟦PII:PERSON:" in red

    restored, stats = hooks.after_llm(f"结果:{red}", conversation_id="conv-1")
    assert stats["n_unrestored"] == 0
    assert "110101199003077890" in restored

    recs = [json.loads(l) for l in
            (hook_env / "hook_log.jsonl").read_text(encoding="utf-8").splitlines()]
    assert [r["event"] for r in recs] == ["before_llm", "after_llm"]
    assert recs[0]["raw"] == text and recs[0]["transformed"] == red
    assert recs[1]["raw"].startswith("结果:") and recs[1]["n_restored"] >= 1
    assert recs[0]["entities"]  # 实体类型统计非空


def test_hook_session_isolation_via_conversation_id(hook_env):
    red = hooks.before_llm("李明的手机号是13800138000", conversation_id="conv-A")
    # 不同 conversation_id 无映射, 占位符保留
    other, stats = hooks.after_llm(red, conversation_id="conv-B")
    assert other == red and stats["n_restored"] == 0


# --------------------------------------------------------------------------- #
# dsh_sim 模拟器
# --------------------------------------------------------------------------- #
def test_dsh_sim_end_to_end(hook_env):
    import dsh_sim
    log_out = hook_env / "sim.jsonl"
    turns = ["张伟的身份证是110101199003077890,帮我查",
             "再查李明, 手机 13800138000"]
    rc = dsh_sim.main(["--turn", turns[0], "--turn", turns[1],
                       "--conversation-id", "test-conv", "--log-out", str(log_out)])
    assert rc == 0
    recs = [json.loads(l) for l in log_out.read_text(encoding="utf-8").splitlines()]
    stages = [(r["turn_no"], r["stage"]) for r in recs]
    assert stages == [(1, "outgoing"), (1, "incoming"), (2, "outgoing"), (2, "incoming")]
    # outgoing 记录不含明文 PII
    for r in recs:
        if r["stage"] == "outgoing":
            for pii in ["110101199003077890", "13800138000", "张伟", "李明"]:
                assert pii not in r["sent_to_llm"]
    # incoming 记录占位符清零
    assert all(r["placeholders_left"] == 0 for r in recs if r["stage"] == "incoming")


# --------------------------------------------------------------------------- #
# standalone_check 真实文档检测
# --------------------------------------------------------------------------- #
def test_standalone_check_docx_roundtrip(hook_env, tmp_path):
    import docx
    import standalone_check

    src = tmp_path / "t.docx"
    d = docx.Document()
    d.add_paragraph("客户张伟, 身份证 110101199003077890。")
    d.add_paragraph("无敏感信息的普通段落。")
    d.save(str(src))

    report_path = tmp_path / "report.json"
    rc = standalone_check.main([str(src), "--report", str(report_path)])
    assert rc == 0
    report = json.loads(report_path.read_text(encoding="utf-8"))
    assert report["all_pass"] is True
    r = report["results"][0]
    assert r["n_replaced"] == 2
    assert r["leaked_pii"] == []
    assert all(u["roundtrip_equal"] and u["no_placeholder_left"]
               for u in r["unit_results"])


def test_standalone_check_missing_file(hook_env, tmp_path):
    import standalone_check
    report_path = tmp_path / "report.json"
    rc = standalone_check.main([str(tmp_path / "nope.txt"),
                                "--report", str(report_path)])
    assert rc == 1


# --------------------------------------------------------------------------- #
# CLI: 默认输出目录 / CSV 凭据 / -d / -v
# --------------------------------------------------------------------------- #
def _make_src(tmp_path):
    src = tmp_path / "报告.txt"
    src.write_text("张伟的身份证是110101199003077890\n普通段落。\n", encoding="utf-8")
    return src


def test_cli_redact_writes_mapping_and_session_csv(hook_env, tmp_path):
    from presidio_pii import cli

    src = _make_src(tmp_path)
    assert cli.main([str(src)]) == 0          # 裸路径等价于 redact

    out = tmp_path / "报告-redacted" / "报告.txt"
    assert out.exists()
    assert "张伟" not in out.read_text(encoding="utf-8")

    mapping_path = out.with_suffix(out.suffix + ".mapping.csv")
    with mapping_path.open(encoding="utf-8-sig", newline="") as f:
        rows = list(csv.DictReader(f))
    assert [r["ptype"] for r in rows] == ["PERSON", "ID_NUMBER"]
    assert rows[0]["original"] == "张伟"
    assert rows[0]["placeholder"] == "⟦PII:PERSON:0001⟧"

    meta_path = out.with_suffix(out.suffix + ".session.csv")
    with meta_path.open(encoding="utf-8-sig", newline="") as f:
        meta = {row["key"]: row["value"] for row in csv.DictReader(f)}
    assert meta["doc_id"] == "报告.txt" and meta["session_id"]


def test_cli_restore_roundtrip_from_csv_without_db(hook_env, tmp_path):
    """默认不落映射库, 仅凭 mapping.csv 也能无损还原。"""
    from presidio_pii import cli

    src = _make_src(tmp_path)
    original = src.read_text(encoding="utf-8")
    assert cli.main([str(src)]) == 0
    assert not (tmp_path / "sessions.db").exists()   # 默认走内存映射库

    redacted = tmp_path / "报告-redacted" / "报告.txt"
    assert cli.main(["restore", str(redacted)]) == 0

    restored = tmp_path / "报告-restored" / "报告.txt"
    assert restored.exists()                  # 同级目录, 不嵌套进 -redacted
    assert restored.read_text(encoding="utf-8") == original


def test_cli_multi_unit_restore_has_no_false_warning(hook_env, tmp_path, capsys):
    """多单元文档: 未命中映射按整篇统计, 不得对每个单元重复报未命中。"""
    import docx
    from presidio_pii import cli

    src = tmp_path / "多段.docx"
    d = docx.Document()
    d.add_paragraph("张伟的身份证是110101199003077890。")
    d.add_paragraph("李明电话 13800138000。")
    d.save(str(src))
    original = [p.text for p in docx.Document(str(src)).paragraphs]

    assert cli.main([str(src)]) == 0
    capsys.readouterr()

    assert cli.main(["restore", str(tmp_path / "多段-redacted" / "多段.docx")]) == 0
    printed = capsys.readouterr().out
    assert "警告" not in printed
    restored = docx.Document(str(tmp_path / "多段-restored" / "多段.docx"))
    assert [p.text for p in restored.paragraphs] == original


def test_cli_out_dir_and_verbose(hook_env, tmp_path, capsys):
    from presidio_pii import cli

    src = _make_src(tmp_path)
    target = tmp_path / "custom"
    assert cli.main(["-v", str(src), "-d", str(target)]) == 0

    assert (target / "报告.txt").exists()
    assert (target / "报告.txt.session.csv").exists()
    printed = capsys.readouterr().out
    assert "[verbose]" in printed and "[2/4]" in printed


def test_cli_output_and_out_dir_are_exclusive(hook_env, tmp_path):
    from presidio_pii import cli

    src = _make_src(tmp_path)
    with pytest.raises(SystemExit):
        cli.main(["redact", str(src), "-o", str(tmp_path / "a.txt"), "-d", str(tmp_path / "b")])


def test_cli_keep_and_whitelist_file(hook_env, tmp_path):
    from presidio_pii import cli

    src = tmp_path / "报告.txt"
    src.write_text("张伟和李明的手机号 13800138000\n", encoding="utf-8")
    wl = tmp_path / "keep.txt"
    wl.write_text("# 保留产品名\n李明\n", encoding="utf-8")

    assert cli.main([str(src), "--whitelist", str(wl), "--keep", "张伟"]) == 0
    out = (tmp_path / "报告-redacted" / "报告.txt").read_text(encoding="utf-8")
    assert "张伟" in out and "李明" in out      # 两个白名单词都保留
    assert "13800138000" not in out             # 其余照常脱敏

    with (tmp_path / "报告-redacted" / "报告.txt.mapping.csv").open(
            encoding="utf-8-sig", newline="") as f:
        rows = list(csv.DictReader(f))
    assert [r["ptype"] for r in rows] == ["PHONE_NUMBER"]


def test_hooks_whitelist_env(hook_env, tmp_path):
    from presidio_pii import hooks

    wl = tmp_path / "wl.txt"
    wl.write_text("张伟\n", encoding="utf-8")
    os.environ["PII_DSH_WHITELIST"] = str(wl)
    hooks.reset_singleton()
    try:
        red = hooks.before_llm("张伟和李明的手机 13800138000", conversation_id="wl-conv")
        assert "张伟" in red
        assert "李明" not in red and "13800138000" not in red
    finally:
        os.environ.pop("PII_DSH_WHITELIST", None)
        hooks.reset_singleton()


# --------------------------------------------------------------------------- #
# 文档提取范围: 表格 / 页眉 / pptx 表格
# --------------------------------------------------------------------------- #
def test_docx_table_and_header_roundtrip(hook_env, tmp_path):
    import docx
    from presidio_pii import cli

    src = tmp_path / "表格.docx"
    d = docx.Document()
    d.add_paragraph("正文段落")
    t = d.add_table(rows=2, cols=2)
    t.cell(0, 0).text = "姓名"
    t.cell(0, 1).text = "身份证号"
    t.cell(1, 0).text = "李芳"
    t.cell(1, 1).text = "110101199003077890"
    d.sections[0].header.paragraphs[0].text = "合肥市数据资源局"
    d.save(str(src))

    assert cli.main([str(src)]) == 0
    red = docx.Document(str(tmp_path / "表格-redacted" / "表格.docx"))
    assert "⟦PII:ID_NUMBER:" in red.tables[0].cell(1, 1).text
    assert "⟦PII:PERSON:" in red.tables[0].cell(1, 0).text
    assert "⟦PII:ORGANIZATION:" in red.sections[0].header.paragraphs[0].text

    assert cli.main(["restore", str(tmp_path / "表格-redacted" / "表格.docx")]) == 0
    back = docx.Document(str(tmp_path / "表格-restored" / "表格.docx"))
    assert back.tables[0].cell(1, 0).text == "李芳"
    assert back.tables[0].cell(1, 1).text == "110101199003077890"
    assert back.sections[0].header.paragraphs[0].text == "合肥市数据资源局"


def test_pptx_table_is_redacted(hook_env, tmp_path):
    from pptx import Presentation
    from pptx.util import Inches
    from presidio_pii import cli

    src = tmp_path / "汇报.pptx"
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[5])
    frame = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(4), Inches(1)).text_frame
    frame.text = "负责人：王建国"
    table = slide.shapes.add_table(2, 2, Inches(1), Inches(3), Inches(4), Inches(1)).table
    table.cell(0, 0).text = "单位"
    table.cell(0, 1).text = "电话"
    table.cell(1, 0).text = "合肥市数据资源局"
    table.cell(1, 1).text = "13900139000"
    prs.save(str(src))

    assert cli.main([str(src)]) == 0
    red = Presentation(str(tmp_path / "汇报-redacted" / "汇报.pptx"))
    texts = []
    for s in red.slides:
        for shape in s.shapes:
            if shape.has_text_frame:
                texts.append(shape.text_frame.text)
            if shape.has_table:
                for row in shape.table.rows:
                    for cell in row.cells:
                        texts.append(cell.text)
    joined = "\n".join(texts)
    assert "王建国" not in joined and "13900139000" not in joined
    assert "⟦PII:ORGANIZATION:" in joined
