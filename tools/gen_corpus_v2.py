"""生成多类型合成测试语料 (v2)。

覆盖 7 类文档: 简历 / 合同 / 公文 / 报告 / 文献 / 工资表 / 通讯录。
每类默认 12 份, 格式按类型取自然载体 (docx / xlsx / pptx / md)。
所有内容为虚构数据; 身份证校验位有效、银行卡通过 Luhn。

产出:
  <out>/<type>/<type>-NN.<ext>
  <out>/manifest.json     每份文档的类型、格式与真值实体
  <out>/whitelist.txt     白名单示例

用法:
  python tools/gen_corpus_v2.py --out testdata_v2 --per-type 12 --seed 42
"""
from __future__ import annotations

import argparse
import json
import random
import sys
from pathlib import Path

import docx
from openpyxl import Workbook
from pptx import Presentation
from pptx.util import Inches

# --------------------------------------------------------------------------- #
# 值池 (全部虚构)
# --------------------------------------------------------------------------- #
PERSONS = [
    "王建国", "李芳", "张伟", "刘志强", "陈晓东", "杨静", "赵磊", "黄敏",
    "周涛", "吴婷", "徐鹏", "孙丽", "马晓东", "朱海波", "胡文华", "郭强",
    "何静怡", "高鹏", "林峰", "罗雪梅", "郑晓燕", "谢文博", "韩雪", "曹阳",
    "邓丽华", "冯志远", "曾庆国", "彭思远", "蒋文静", "蔡明辉", "潘俊杰",
    "杜雪峰", "叶文倩", "程海燕", "苏浩然", "魏明轩", "吕晓雯", "丁一凡",
    "任建军", "沈佳怡", "姚志刚", "卢晓东", "姜文涛", "崔晓峰", "钟丽娜",
    "谭俊杰", "陆文斌", "汪海燕", "范志强", "金晓明", "石文杰", "廖海涛",
    "贾晓丽", "夏文静", "韦志远", "付晓东", "方文娟", "白晓峰", "邹海燕",
    "孟志强",
]

NATIONAL_ORGS = [
    "国家发展和改革委员会", "工业和信息化部", "国家互联网信息办公室", "财政部",
    "教育部", "国家数据局", "国务院办公厅", "公安部", "国家统计局", "国家税务总局",
]
PROVINCE_ORGS = [
    "安徽省数据资源管理局", "安徽省发展和改革委员会", "江苏省工业和信息化厅",
    "浙江省财政厅", "广东省政务服务数据管理局", "山东省大数据局",
    "河南省发展和改革委员会", "四川省经济和信息化厅", "湖北省财政厅",
    "福建省数字福建建设领导小组办公室", "河北省数据和政务服务局", "湖南省工业和信息化厅",
]
CITY_ORGS = [
    "合肥市数据资源局", "南京市大数据管理局", "苏州市行政审批局", "杭州市数据资源管理局",
    "芜湖市信息中心", "宁波市大数据发展管理局", "武汉市政务服务和大数据管理局",
    "成都市大数据中心", "郑州市大数据管理局", "长沙市数据资源管理局",
    "济南市大数据局", "福州市大数据发展管理委员会",
]
COUNTY_ORGS = [
    "肥东县数据资源局", "长丰县财政局", "庐江县发展和改革委员会", "无为市信息中心",
    "巢湖市政务服务管理局", "桐城市数据资源局", "肥西县数据资源局", "当涂县财政局",
    "溧阳市大数据管理局", "昆山市行政审批局", "慈溪市数据资源中心", "余姚市信息中心",
]
COMPANIES = [
    "安徽智数科技有限公司", "合肥云启信息技术有限公司", "中科云谷（合肥）数据科技有限公司",
    "皖信通网络科技有限公司", "安徽数安信息技术股份有限公司", "合肥联创软件有限公司",
    "江苏鸿信数据服务有限公司", "南京数联智能科技有限公司", "浙江云枢信息技术有限公司",
    "杭州智联数据服务有限公司", "山东齐鲁数字科技有限公司", "河南中原数据服务有限公司",
    "湖北楚天信息技术有限公司", "四川天府数智科技有限公司",
]
ORGS = NATIONAL_ORGS + PROVINCE_ORGS + CITY_ORGS + COUNTY_ORGS + COMPANIES

ABBR = {
    "安徽智数科技有限公司": "智数科技",
    "合肥云启信息技术有限公司": "云启信息",
    "中科云谷（合肥）数据科技有限公司": "中科云谷",
    "皖信通网络科技有限公司": "皖信通",
    "安徽数安信息技术股份有限公司": "数安信息",
    "合肥联创软件有限公司": "联创软件",
    "江苏鸿信数据服务有限公司": "鸿信数据",
    "南京数联智能科技有限公司": "数联智能",
    "浙江云枢信息技术有限公司": "云枢信息",
    "杭州智联数据服务有限公司": "智联数据",
    "山东齐鲁数字科技有限公司": "齐鲁数字",
    "河南中原数据服务有限公司": "中原数据",
    "湖北楚天信息技术有限公司": "楚天信息",
    "四川天府数智科技有限公司": "天府数智",
    "国家数据局": "数据局",
    "安徽省数据资源管理局": "省数据局",
    "合肥市数据资源局": "市数据局",
    "国家发展和改革委员会": "国家发改委",
    "工业和信息化部": "工信部",
    "国家互联网信息办公室": "国家网信办",
}

ADDRESSES = [
    "安徽省合肥市蜀山区黄山路600号", "合肥市包河区中山路1号", "合肥市庐阳区长江中路88号",
    "合肥市瑶海区临泉路200号", "合肥市政务区潜山路100号",
    "江苏省南京市鼓楼区北京西路70号", "南京市玄武区中山路300号", "南京市建邺区江东中路50号",
    "南京市秦淮区中华路20号", "南京市栖霞区仙林大道10号",
    "浙江省杭州市西湖区文三路90号", "杭州市滨江区江南大道100号", "杭州市拱墅区莫干山路200号",
    "杭州市余杭区文一西路969号", "杭州市上城区解放路50号",
    "芜湖市镜湖区文化路8号", "芜湖市弋江区九华南路10号", "芜湖市鸠江区政通路20号",
    "芜湖市湾沚区中山路30号", "芜湖市繁昌区长江南路40号",
    "成都市武侯区天府大道北段1700号", "成都市锦江区红星路三段1号", "成都市青羊区人民中路二段10号",
    "成都市高新区天府五街200号", "成都市金牛区蜀汉路100号",
    "武汉市江岸区沿江大道188号", "武汉市武昌区中南路7号", "武汉市洪山区珞喻路100号",
    "武汉市江汉区新华路25号", "武汉市汉阳区琴台大道50号",
    "郑州市金水区农业路60号", "郑州市郑东新区商务外环路22号", "郑州市中原区建设西路100号",
    "郑州市二七区大学路80号", "郑州市管城区商城路15号",
    "长沙市岳麓区麓谷大道658号", "长沙市芙蓉区五一大道100号", "长沙市天心区湘府路200号",
    "长沙市开福区芙蓉北路30号", "长沙市雨花区万家丽中路50号",
    "济南市历下区经十路100号", "济南市市中区纬二路50号", "济南市槐荫区经四路200号",
    "济南市天桥区无影山路30号", "济南市历城区华龙路10号",
    "福州市鼓楼区五四路100号", "福州市台江区江滨中大道50号", "福州市仓山区金山大道200号",
    "福州市晋安区福新路30号", "福州市马尾区江滨东大道10号",
]

TITLES = [
    "局长", "副局长", "处长", "副处长", "科长", "主任", "副主任", "总经理", "副总经理",
    "董事长", "党委书记", "总工程师", "工程师", "高级工程师", "教授", "研究员", "院长",
    "校长", "所长", "队长", "经理", "秘书长", "技术总监", "财务总监", "项目经理",
    "架构师", "分析师", "顾问", "专员", "主管",
]
DEPARTMENTS = [
    "办公室", "综合处", "财务处", "人事处", "技术部", "研发部", "数据资源科", "信息中心",
    "运维部", "市场部", "法务部", "采购部", "审计处", "监察室", "项目办", "规划处",
    "政策法规处", "政务服务科", "网络科", "安全科", "应用推广科", "标准规范科",
    "档案室", "后勤保障部", "客户服务部",
]
EDUCATION = ["本科", "硕士", "博士", "大专", "中专"]
SCHOOLS = [
    "北京大学", "清华大学", "中国科学技术大学", "复旦大学", "上海交通大学", "浙江大学",
    "南京大学", "武汉大学", "中山大学", "四川大学", "合肥工业大学", "安徽大学",
    "南京理工大学", "东南大学", "华中科技大学", "西安交通大学", "哈尔滨工业大学",
    "同济大学", "厦门大学", "山东大学",
]
MAJORS = [
    "计算机科学与技术", "软件工程", "信息安全", "数据科学与大数据技术", "信息管理与信息系统",
    "电子信息工程", "通信工程", "自动化", "统计学", "数学与应用数学", "工商管理", "会计学",
    "法学", "行政管理", "汉语言文学", "网络工程", "人工智能", "物联网工程", "电子商务",
    "公共事业管理",
]
CERTIFICATES = [
    "系统架构设计师", "信息系统项目管理师", "数据库系统工程师", "网络工程师", "软件设计师",
    "注册信息安全工程师", "CISSP", "CISP", "PMP", "注册会计师", "法律职业资格",
    "一级建造师", "人力资源管理师", "中级经济师", "高级工程师",
]
CONTRACT_SUBJECTS = [
    "软件开发", "系统集成", "数据治理", "运维服务", "网络安全评估", "机房托管",
    "云资源租用", "技术咨询", "人员外包", "设备采购", "平台运营", "等级保护测评",
    "密码应用改造", "数据清洗", "接口开发",
]
TERMS = ["12个月", "18个月", "24个月", "36个月", "6个月", "一年", "两年", "三年",
         "90日历天", "180日历天"]
SECURITY_LEVELS = ["公开", "内部", "秘密", "机密"]
URGENCIES = ["特急", "加急", "平急", "普通"]
DOC_PREFIXES = ["皖数据", "皖发改", "苏工信", "浙财", "粤政数", "鲁数", "豫发改", "川经信",
                "鄂财", "闽数"]
JOURNALS = [
    "计算机学报", "软件学报", "中国科学：信息科学", "电子学报", "通信学报",
    "计算机研究与发展", "自动化学报", "中文信息学报", "数据分析与知识发现", "情报学报",
    "管理世界", "中国行政管理", "电子政务", "信息安全学报", "网络与信息安全学报",
]
FUNDS = [
    "国家自然科学基金（62072123）", "国家重点研发计划（2023YFB3204500）",
    "国家社会科学基金（22BGL123）", "安徽省重点研究与开发计划（202304a05020001）",
    "江苏省自然科学基金（BK20230123）", "浙江省科技计划项目（2023C01234）",
    "广东省基础与应用基础研究基金（2023A1515011234）", "山东省重点研发计划（2023JZZX0101）",
    "河南省科技攻关项目（232102210123）", "四川省科技计划项目（2023YFS0123）",
    "湖北省自然科学基金（2023CFB123）", "福建省科技计划项目（2023H0012）",
    "国家密码管理局密码应用专项（2023-MM-012）", "教育部人文社科项目（23YJC870012）",
    "中国博士后科学基金（2023M731234）",
]

_WEIGHTS = [7, 9, 10, 5, 8, 4, 2, 1, 6, 3, 5, 7, 9, 10, 5, 8, 4, 2, 1]
_MOD_MAP = "10X98765432"
_ID_PREFIXES = ["340102", "320106", "330106", "340202", "340122", "340111",
                "510107", "420102", "410105", "430104", "370102", "350102"]
_CREDIT_CHARS = "0123456789ABCDEFGHJKLMNPQRTUWXY"
_DOMAINS = ["example.com", "example.org", "example.net"]
_AREA_CODES = ["0551", "025", "0571", "0553", "028", "027", "0371", "0731", "0531", "0591"]


# --------------------------------------------------------------------------- #
# 生成工具
# --------------------------------------------------------------------------- #
def make_id(rng: random.Random) -> str:
    body = (rng.choice(_ID_PREFIXES) + f"{rng.randint(1975, 1998):04d}"
            + f"{rng.randint(1, 12):02d}{rng.randint(1, 28):02d}"
            + f"{rng.randint(1, 999):03d}")
    total = sum(int(body[i]) * _WEIGHTS[i] for i in range(17))
    return body + _MOD_MAP[total % 11]


def make_phone(rng: random.Random) -> str:
    return rng.choice(["138", "139", "150", "158", "186", "177"]) + "".join(
        str(rng.randint(0, 9)) for _ in range(8))


def make_office_phone(rng: random.Random) -> str:
    return f"{rng.choice(_AREA_CODES)}-{rng.randint(60000000, 69999999)}"


def make_email(rng: random.Random) -> str:
    return f"user{rng.randint(10000, 99999)}@{rng.choice(_DOMAINS)}"


def make_credit(rng: random.Random) -> str:
    return "91340100" + "".join(rng.choice(_CREDIT_CHARS) for _ in range(10))


def make_bank_card(rng: random.Random) -> str:
    """16 位银行卡号, 通过标准 Luhn 校验。

    标准 Luhn: 从右往左, 校验位为第 1 位 (不乘 2), 其余隔位乘 2。
    """
    digits = [int(c) for c in rng.choice(["622202", "621226", "622848", "621700", "436742"])]
    digits += [rng.randint(0, 9) for _ in range(9)]

    def luhn(ds):
        total = 0
        for i, d in enumerate(reversed(ds)):
            if i % 2 == 1:          # 校验位是第 0 位, 不参与乘 2
                d *= 2
                if d > 9:
                    d -= 9
            total += d
        return total

    for check in range(10):
        if luhn(digits + [check]) % 10 == 0:
            return "".join(map(str, digits + [check]))
    return "".join(map(str, digits + [0]))


def make_amount(rng: random.Random) -> str:
    return f"{rng.randint(5, 2000)}.{rng.randint(0, 99):02d}"


class Truth:
    """真值收集器。"""

    def __init__(self):
        self.items = []

    def add(self, text: str, etype: str, field: str, hard: bool = False) -> str:
        self.items.append({"text": text, "type": etype, "field": field, "hard": hard})
        return text


# --------------------------------------------------------------------------- #
# 各类型文档
# --------------------------------------------------------------------------- #
def build_resume(rng, path: Path, truth: Truth) -> None:
    name = rng.choice(PERSONS)
    phone = make_phone(rng)
    email = make_email(rng)
    id_no = make_id(rng)
    address = rng.choice(ADDRESSES)
    company = rng.choice(COMPANIES)
    title = rng.choice(TITLES)
    school = rng.choice(SCHOOLS)
    major = rng.choice(MAJORS)
    edu = rng.choice(EDUCATION)
    cert = rng.choice(CERTIFICATES)
    contact = rng.choice(PERSONS)

    truth.add(name, "PERSON", "姓名")
    truth.add(phone, "PHONE_NUMBER", "联系电话")
    truth.add(email, "EMAIL_ADDRESS", "邮箱")
    truth.add(id_no, "ID_NUMBER", "身份证号")
    truth.add(address, "LOCATION", "现居地址")
    truth.add(company, "ORGANIZATION", "工作单位")
    truth.add(ABBR.get(company, company), "ORGANIZATION", "单位简称", hard=True)
    truth.add(title, "JOB_TITLE", "职位")
    truth.add(contact, "PERSON", "紧急联系人")

    d = docx.Document()
    d.add_heading("个人简历", level=0)
    d.add_paragraph(f"姓名：{name}    性别：{'男' if rng.random() < 0.6 else '女'}"
                    f"    出生年月：{rng.randint(1985, 2000)}-{rng.randint(1, 12):02d}")
    d.add_paragraph(f"联系电话：{phone}    电子邮箱：{email}")
    d.add_paragraph(f"身份证号：{id_no}")
    d.add_paragraph(f"现居地址：{address}")
    d.add_paragraph(f"籍贯：{rng.choice(ADDRESSES)[:6]}")

    d.add_heading("教育经历", level=1)
    table = d.add_table(rows=2, cols=4)
    for ci, head in enumerate(["起止时间", "学校", "专业", "学历"]):
        table.cell(0, ci).text = head
    table.cell(1, 0).text = f"{rng.randint(2006, 2012)}.09-{rng.randint(2013, 2016)}.06"
    table.cell(1, 1).text = school
    table.cell(1, 2).text = major
    table.cell(1, 3).text = edu

    d.add_heading("工作经历", level=1)
    company_line = f"{rng.randint(2014, 2020)}.07-至今    {company}    {title}"
    abbr = ABBR.get(company)
    if abbr and abbr != company:
        company_line += f"（以下简称“{abbr}”）"
    d.add_paragraph(company_line)
    d.add_paragraph(f"负责{rng.choice(CONTRACT_SUBJECTS)}相关工作，参与项目{ rng.randint(3, 12) }个。")

    d.add_heading("项目经历", level=1)
    d.add_paragraph(f"{rng.randint(2020, 2024)}年    {rng.choice(CITY_ORGS)}"
                    f"{rng.choice(CONTRACT_SUBJECTS)}项目，担任{title}。")

    d.add_heading("技能证书", level=1)
    d.add_paragraph(f"• {cert}")
    d.add_paragraph(f"• {rng.choice(CERTIFICATES)}")
    d.add_paragraph(f"• 熟练使用 {rng.choice(['Java', 'Python', 'Go', 'SQL'])} 与 "
                    f"{rng.choice(['Spring Cloud', 'Kubernetes', 'Spark', 'Flink'])}。")

    d.add_heading("其他", level=1)
    d.add_paragraph(f"紧急联系人：{contact}    关系：{rng.choice(['配偶', '父母', '兄弟'])}")
    d.add_paragraph("自我评价：工作认真负责，具备良好的沟通与团队协作能力。")
    d.save(str(path))


def build_contract(rng, path: Path, truth: Truth) -> None:
    party_a = rng.choice(ORGS)
    party_b = rng.choice(COMPANIES)
    person_a = rng.choice(PERSONS)
    person_b = rng.choice(PERSONS)
    addr_a = rng.choice(ADDRESSES)
    addr_b = rng.choice(ADDRESSES)
    phone_a = make_phone(rng)
    phone_b = make_phone(rng)
    credit_a = make_credit(rng)
    credit_b = make_credit(rng)
    bank = make_bank_card(rng)
    email = make_email(rng)
    subject = rng.choice(CONTRACT_SUBJECTS)
    amount = make_amount(rng)
    term = rng.choice(TERMS)

    truth.add(party_a, "ORGANIZATION", "甲方")
    truth.add(party_b, "ORGANIZATION", "乙方")
    truth.add(person_a, "PERSON", "甲方法定代表人")
    truth.add(person_b, "PERSON", "乙方法定代表人")
    truth.add(addr_a, "LOCATION", "甲方地址")
    truth.add(addr_b, "LOCATION", "乙方地址")
    truth.add(phone_a, "PHONE_NUMBER", "甲方电话")
    truth.add(phone_b, "PHONE_NUMBER", "乙方电话")
    truth.add(credit_a, "CREDIT_CODE", "甲方信用代码")
    truth.add(credit_b, "CREDIT_CODE", "乙方信用代码")
    truth.add(bank, "BANK_CARD", "收款账号")
    truth.add(email, "EMAIL_ADDRESS", "联系邮箱")
    truth.add(ABBR.get(party_b, party_b), "ORGANIZATION", "乙方简称", hard=True)

    d = docx.Document()
    d.add_heading(f"{subject}服务合同", level=0)
    d.add_paragraph(f"合同编号：HT-{rng.randint(2023, 2024)}-{rng.randint(1000, 9999)}")
    d.add_paragraph("")
    d.add_paragraph("甲方（全称）：" + party_a)
    d.add_paragraph("统一社会信用代码：" + credit_a)
    d.add_paragraph("法定代表人：" + person_a)
    d.add_paragraph("地址：" + addr_a)
    d.add_paragraph("联系电话：" + phone_a)
    d.add_paragraph("")
    d.add_paragraph("乙方（全称）：" + party_b
                    + (f"（以下简称“{ABBR[party_b]}”）" if party_b in ABBR else ""))
    d.add_paragraph("统一社会信用代码：" + credit_b)
    d.add_paragraph("法定代表人：" + person_b)
    d.add_paragraph("地址：" + addr_b)
    d.add_paragraph("联系电话：" + phone_b)
    d.add_paragraph("电子邮箱：" + email)
    d.add_paragraph("")
    d.add_paragraph("第一条 项目内容")
    d.add_paragraph(f"乙方向甲方提供{subject}服务，具体范围以附件《服务清单》为准。")
    d.add_paragraph("第二条 合同金额")
    d.add_paragraph(f"本合同总金额为人民币 {amount} 万元（含税）。")
    d.add_paragraph("第三条 履行期限")
    d.add_paragraph(f"自合同签订之日起 {term}。")
    d.add_paragraph("第四条 付款方式")
    table = d.add_table(rows=4, cols=3)
    for ci, head in enumerate(["阶段", "比例", "金额（万元）"]):
        table.cell(0, ci).text = head
    for ri, (stage, ratio) in enumerate([("合同签订", "30%"), ("中期验收", "40%"), ("终验合格", "30%")], 1):
        table.cell(ri, 0).text = stage
        table.cell(ri, 1).text = ratio
        table.cell(ri, 2).text = f"{float(amount) * int(ratio[:-1]) / 100:.2f}"
    d.add_paragraph("第五条 收款账户")
    d.add_paragraph("开户银行：中国工商银行" + rng.choice(["合肥分行", "南京分行", "杭州分行"]))
    d.add_paragraph("银行账号：" + bank)
    d.add_paragraph("第六条 违约责任")
    d.add_paragraph("任何一方违约，应向对方支付合同总额 5% 的违约金。")
    d.add_paragraph("第七条 争议解决")
    d.add_paragraph("因本合同产生的争议，提交" + rng.choice(CITY_ORGS) + "所在地人民法院管辖。")
    d.add_paragraph("")
    d.add_paragraph("甲方（盖章）：" + party_a + "        乙方（盖章）：" + party_b)
    d.add_paragraph(f"法定代表人：{person_a}              法定代表人：{person_b}")
    d.add_paragraph(f"{rng.randint(2023, 2024)}年{rng.randint(1, 12)}月{rng.randint(1, 28)}日")
    d.save(str(path))


def build_official(rng, path: Path, truth: Truth) -> None:
    issuer = rng.choice(PROVINCE_ORGS + CITY_ORGS)
    receiver = rng.choice(NATIONAL_ORGS + CITY_ORGS)
    cc = rng.choice(CITY_ORGS + COUNTY_ORGS)
    signer = rng.choice(PERSONS)
    contact = rng.choice(PERSONS)
    phone = make_office_phone(rng)
    prefix = rng.choice(DOC_PREFIXES)
    doc_no = f"{prefix}〔{rng.randint(2022, 2024)}〕{rng.randint(1, 200)}号"
    level = rng.choice(SECURITY_LEVELS)
    urgency = rng.choice(URGENCIES)
    subject = rng.choice(CONTRACT_SUBJECTS)
    year, month, day = rng.randint(2022, 2024), rng.randint(1, 12), rng.randint(1, 28)

    truth.add(issuer, "ORGANIZATION", "发文机关")
    truth.add(receiver, "ORGANIZATION", "主送机关")
    truth.add(cc, "ORGANIZATION", "抄送机关")
    truth.add(signer, "PERSON", "签发人")
    truth.add(contact, "PERSON", "联系人")
    truth.add(phone, "PHONE_NUMBER", "联系电话")
    truth.add(ABBR.get(issuer, issuer), "ORGANIZATION", "发文机关简称", hard=True)

    d = docx.Document()
    d.add_paragraph(issuer)
    if issuer in ABBR:
        d.add_paragraph(f"（以下简称“{ABBR[issuer]}”）")
    d.add_paragraph(doc_no)
    d.add_paragraph(f"密级：{level}    紧急程度：{urgency}")
    d.add_paragraph("")
    d.add_heading(f"关于推进{subject}工作的通知", level=1)
    d.add_paragraph(receiver + "：")
    d.add_paragraph(f"为贯彻落实{receiver}关于{subject}的工作部署，进一步提升数据资源管理水平，"
                    f"现就有关事项通知如下。")
    d.add_paragraph("一、总体要求")
    d.add_paragraph("坚持统筹规划、需求导向、安全可控的原则，加快推进相关工作。")
    d.add_paragraph("二、重点任务")
    d.add_paragraph(f"（一）完善{subject}工作机制，明确责任分工。")
    d.add_paragraph("（二）加强数据质量管理，建立常态化核查机制。")
    d.add_paragraph("（三）强化安全保障，落实等级保护要求。")
    d.add_paragraph("三、工作要求")
    d.add_paragraph("各单位要高度重视，明确专人负责，确保按期完成。")
    d.add_paragraph("")
    d.add_paragraph(f"联系人：{contact}    联系电话：{phone}")
    d.add_paragraph("")
    d.add_paragraph(issuer)
    d.add_paragraph(f"{year}年{month}月{day}日")
    d.add_paragraph(f"抄送：{cc}。")
    d.add_paragraph(f"签发人：{signer}")
    d.save(str(path))


def build_report_docx(rng, path: Path, truth: Truth) -> None:
    builder = rng.choice(COMPANIES)
    owner = rng.choice(ORGS)
    leader = rng.choice(PERSONS)
    contact = rng.choice(PERSONS)
    phone = make_phone(rng)
    email = make_email(rng)
    address = rng.choice(ADDRESSES)
    members = rng.sample(PERSONS, 3)
    experts = rng.sample(PERSONS, 2)
    budget = make_amount(rng)

    for m in members + experts:
        truth.add(m, "PERSON", "项目成员/专家")
    truth.add(leader, "PERSON", "项目负责人")
    truth.add(contact, "PERSON", "联系人")
    truth.add(builder, "ORGANIZATION", "编制单位")
    truth.add(owner, "ORGANIZATION", "建设单位")
    truth.add(phone, "PHONE_NUMBER", "联系电话")
    truth.add(email, "EMAIL_ADDRESS", "邮箱")
    truth.add(address, "LOCATION", "单位地址")

    d = docx.Document()
    d.add_heading("项目建议书", level=0)
    d.add_paragraph(f"项目名称：{rng.choice(CONTRACT_SUBJECTS)}平台建设项目")
    d.add_paragraph(f"建设单位：{owner}")
    d.add_paragraph(f"编制单位：{builder}")
    d.add_paragraph(f"项目负责人：{leader}")
    d.add_paragraph(f"联系电话：{phone}    电子邮箱：{email}")
    d.add_paragraph(f"单位地址：{address}")
    d.add_paragraph("")
    d.add_heading("第一章 项目概述", level=1)
    d.add_paragraph("本项目围绕数据资源整合共享，建设统一的管理平台。")
    d.add_heading("第二章 建设内容", level=1)
    d.add_paragraph("包括数据治理、接口开发、门户建设、运维保障等内容。")
    d.add_heading("第三章 投资估算", level=1)
    d.add_paragraph(f"项目总投资约 {budget} 万元。")
    d.add_heading("第四章 项目组", level=1)
    table = d.add_table(rows=len(members) + 1, cols=3)
    for ci, head in enumerate(["姓名", "角色", "联系电话"]):
        table.cell(0, ci).text = head
    for ri, m in enumerate(members, 1):
        role = rng.choice(["项目经理", "技术负责人", "实施工程师"])
        table.cell(ri, 0).text = m
        table.cell(ri, 1).text = role
        table.cell(ri, 2).text = make_phone(rng)
        truth.add(role, "JOB_TITLE", "角色")
    d.add_heading("第五章 评审专家", level=1)
    d.add_paragraph("、".join(experts) + "等专家参与评审。")
    d.add_paragraph(f"联系单位：{builder}    联系人：{contact}")
    d.save(str(path))


def build_report_pptx(rng, path: Path, truth: Truth) -> None:
    owner = rng.choice(ORGS)
    leader = rng.choice(PERSONS)
    phone = make_phone(rng)
    email = make_email(rng)
    members = rng.sample(PERSONS, 2)

    truth.add(owner, "ORGANIZATION", "建设单位")
    truth.add(leader, "PERSON", "项目负责人")
    truth.add(phone, "PHONE_NUMBER", "联系电话")
    truth.add(email, "EMAIL_ADDRESS", "邮箱")
    for m in members:
        truth.add(m, "PERSON", "项目成员")

    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[0])
    slide.shapes.title.text = "项目建设方案汇报"
    slide.placeholders[1].text = f"{owner}    汇报人：{leader}"

    slide = prs.slides.add_slide(prs.slide_layouts[1])
    slide.shapes.title.text = "项目概况"
    body = slide.placeholders[1].text_frame
    body.text = f"建设单位：{owner}"
    body.add_paragraph().text = f"项目负责人：{leader}"
    body.add_paragraph().text = f"联系电话：{phone}"
    body.add_paragraph().text = f"电子邮箱：{email}"

    slide = prs.slides.add_slide(prs.slide_layouts[5])
    slide.shapes.title.text = "项目组成员"
    table = slide.shapes.add_table(len(members) + 1, 3, Inches(1), Inches(2),
                                   Inches(6), Inches(1.5)).table
    for ci, head in enumerate(["姓名", "角色", "联系电话"]):
        table.cell(0, ci).text = head
    for ri, m in enumerate(members, 1):
        table.cell(ri, 0).text = m
        table.cell(ri, 1).text = "技术负责人"
        table.cell(ri, 2).text = make_phone(rng)
    truth.add("技术负责人", "JOB_TITLE", "角色")

    slide = prs.slides.add_slide(prs.slide_layouts[1])
    slide.shapes.title.text = "下一步计划"
    body = slide.placeholders[1].text_frame
    body.text = "完善需求调研，推进方案评审。"
    prs.save(str(path))


def build_paper_md(rng, path: Path, truth: Truth) -> None:
    authors = rng.sample(PERSONS, 3)
    affiliation = rng.choice(SCHOOLS + PROVINCE_ORGS)
    email = make_email(rng)
    fund = rng.choice(FUNDS)
    journal = rng.choice(JOURNALS)
    ref_authors = rng.sample(PERSONS, 2)

    for a in authors:
        truth.add(a, "PERSON", "作者")
    for a in ref_authors:
        truth.add(a, "PERSON", "参考文献作者")
    truth.add(affiliation, "ORGANIZATION", "作者单位")
    truth.add(email, "EMAIL_ADDRESS", "通讯邮箱")

    text = f"""# 面向政务数据共享的中文实体识别方法研究

{authors[0]}，{authors[1]}，{authors[2]}

（{affiliation}）

**摘要**：针对政务文本中实体边界模糊的问题，提出一种融合词典与统计特征的中文实体识别方法。
**关键词**：中文实体识别；政务数据；边界消歧

## 1 引言

政务数据共享是数字政府建设的基础。已有研究[{ref_authors[0]}, {ref_authors[1]}]表明，
中文实体识别在专业领域仍面临召回与精度难以兼顾的问题。

## 2 方法

本文采用姓名字典与上下文规则相结合的方式，并在{affiliation}的真实语料上验证。

## 3 实验

实验在自建语料上进行，评价指标为准确率、召回率与 F1 值。

## 4 结论

所提方法在政务文本上取得较好效果。

## 基金项目

{fund}

## 参考文献

[1] {ref_authors[0]}. 政务数据治理研究[J]. {journal}, 2023, 45(3): 1-12.
[2] {ref_authors[1]}. 中文实体识别综述[J]. {journal}, 2022, 44(8): 33-45.

通讯作者：{authors[0]}，邮箱：{email}
"""
    path.write_text(text, encoding="utf-8")


def build_paper_docx(rng, path: Path, truth: Truth) -> None:
    authors = rng.sample(PERSONS, 3)
    affiliation = rng.choice(SCHOOLS)
    email = make_email(rng)
    fund = rng.choice(FUNDS)
    journal = rng.choice(JOURNALS)
    ref_authors = rng.sample(PERSONS, 2)

    for a in authors + ref_authors:
        truth.add(a, "PERSON", "作者")
    truth.add(affiliation, "ORGANIZATION", "作者单位")
    truth.add(email, "EMAIL_ADDRESS", "通讯邮箱")

    d = docx.Document()
    d.add_heading("面向政务数据共享的中文实体识别方法研究", level=0)
    d.add_paragraph("，".join(authors))
    d.add_paragraph(f"（{affiliation}）")
    d.add_paragraph("摘要：针对政务文本中实体边界模糊的问题，提出一种融合词典与统计特征的方法。")
    d.add_paragraph("关键词：中文实体识别；政务数据；边界消歧")
    d.add_heading("1 引言", level=1)
    d.add_paragraph("政务数据共享是数字政府建设的基础。")
    d.add_heading("2 方法", level=1)
    d.add_paragraph(f"本文在{affiliation}的真实语料上验证所提方法。")
    d.add_heading("3 实验", level=1)
    d.add_paragraph("实验以准确率、召回率与 F1 值为评价指标。")
    d.add_heading("基金项目", level=1)
    d.add_paragraph(fund)
    d.add_heading("参考文献", level=1)
    d.add_paragraph(f"[1] {ref_authors[0]}. 政务数据治理研究[J]. {journal}, 2023, 45(3): 1-12.")
    d.add_paragraph(f"[2] {ref_authors[1]}. 中文实体识别综述[J]. {journal}, 2022, 44(8): 33-45.")
    d.add_paragraph(f"通讯作者：{authors[0]}，邮箱：{email}")
    d.save(str(path))


def build_payroll(rng, path: Path, truth: Truth) -> None:
    rows = rng.randint(10, 20)
    wb = Workbook()
    ws = wb.active
    ws.title = "工资表"
    headers = ["工号", "姓名", "身份证号", "银行卡号", "部门", "岗位",
               "基本工资", "绩效工资", "社保个人", "公积金", "个税", "实发工资", "联系电话"]
    ws.append(headers)
    for i in range(rows):
        name = rng.choice(PERSONS)
        id_no = make_id(rng)
        bank = make_bank_card(rng)
        phone = make_phone(rng)
        base = rng.randint(4000, 18000)
        perf = rng.randint(0, 8000)
        social = int(base * 0.105)
        fund = int(base * 0.07)
        tax = max(0, int((base + perf - social - fund - 5000) * 0.1))
        net = base + perf - social - fund - tax
        dept = rng.choice(DEPARTMENTS)
        title = rng.choice(TITLES)
        ws.append([f"E{2024000 + i}", name, id_no, bank, dept,
                   title, base, perf, social, fund, tax, net, phone])
        truth.add(name, "PERSON", "姓名")
        truth.add(id_no, "ID_NUMBER", "身份证号")
        truth.add(bank, "BANK_CARD", "银行卡号")
        truth.add(phone, "PHONE_NUMBER", "联系电话")
        truth.add(title, "JOB_TITLE", "岗位")
    wb.save(str(path))


def build_contacts_xlsx(rng, path: Path, truth: Truth) -> None:
    rows = rng.randint(12, 20)
    wb = Workbook()
    ws = wb.active
    ws.title = "通讯录"
    ws.append(["姓名", "部门", "职务", "手机", "办公电话", "邮箱", "地址"])
    for _ in range(rows):
        name = rng.choice(PERSONS)
        mobile = make_phone(rng)
        office = make_office_phone(rng)
        email = make_email(rng)
        address = rng.choice(ADDRESSES)
        title = rng.choice(TITLES)
        ws.append([name, rng.choice(DEPARTMENTS), title, mobile, office, email, address])
        truth.add(name, "PERSON", "姓名")
        truth.add(mobile, "PHONE_NUMBER", "手机")
        truth.add(office, "PHONE_NUMBER", "办公电话")
        truth.add(email, "EMAIL_ADDRESS", "邮箱")
        truth.add(address, "LOCATION", "地址")
        truth.add(title, "JOB_TITLE", "职务")
    wb.save(str(path))


def build_contacts_docx(rng, path: Path, truth: Truth) -> None:
    rows = rng.randint(10, 15)
    d = docx.Document()
    d.add_heading("单位通讯录", level=0)
    d.add_paragraph(f"更新日期：{rng.randint(2023, 2024)}年{rng.randint(1, 12)}月")
    table = d.add_table(rows=rows + 1, cols=6)
    for ci, head in enumerate(["姓名", "部门", "职务", "手机", "办公电话", "邮箱"]):
        table.cell(0, ci).text = head
    for ri in range(1, rows + 1):
        name = rng.choice(PERSONS)
        mobile = make_phone(rng)
        office = make_office_phone(rng)
        email = make_email(rng)
        title = rng.choice(TITLES)
        values = [name, rng.choice(DEPARTMENTS), title, mobile, office, email]
        for ci, value in enumerate(values):
            table.cell(ri, ci).text = value
        truth.add(name, "PERSON", "姓名")
        truth.add(mobile, "PHONE_NUMBER", "手机")
        truth.add(office, "PHONE_NUMBER", "办公电话")
        truth.add(email, "EMAIL_ADDRESS", "邮箱")
        truth.add(title, "JOB_TITLE", "职务")
    d.save(str(path))


# --------------------------------------------------------------------------- #
# 编排
# --------------------------------------------------------------------------- #
BUILDERS = {
    "resume": (build_resume, "docx"),
    "contract": (build_contract, "docx"),
    "official": (build_official, "docx"),
    "report": (build_report_docx, "docx"),
    "report_pptx": (build_report_pptx, "pptx"),
    "paper": (build_paper_md, "md"),
    "paper_docx": (build_paper_docx, "docx"),
    "payroll": (build_payroll, "xlsx"),
    "contacts": (build_contacts_xlsx, "xlsx"),
    "contacts_docx": (build_contacts_docx, "docx"),
}

# 每类 12 份: report 6 docx + 6 pptx, paper 6 md + 6 docx, contacts 6 xlsx + 6 docx
PLAN = {
    "resume": ["resume"] * 12,
    "contract": ["contract"] * 12,
    "official": ["official"] * 12,
    "report": ["report"] * 6 + ["report_pptx"] * 6,
    "paper": ["paper"] * 6 + ["paper_docx"] * 6,
    "payroll": ["payroll"] * 12,
    "contacts": ["contacts"] * 6 + ["contacts_docx"] * 6,
}


def generate(out_dir: Path, per_type: int = 12, seed: int = 42) -> dict:
    rng = random.Random(seed)
    manifest = {"seed": seed, "per_type": per_type, "docs": [], "types": {}}

    for dtype, variants in PLAN.items():
        count = min(per_type, len(variants))
        type_dir = out_dir / dtype
        type_dir.mkdir(parents=True, exist_ok=True)
        for i in range(count):
            kind = variants[i]
            builder, ext = BUILDERS[kind]
            name = f"{dtype}-{i + 1:02d}"
            path = type_dir / f"{name}.{ext}"
            truth = Truth()
            builder(rng, path, truth)
            manifest["docs"].append({
                "file": str(path.relative_to(out_dir)).replace("\\", "/"),
                "type": dtype,
                "format": ext,
                "entities": truth.items,
            })
        manifest["types"][dtype] = count

    manifest["whitelist"] = ["国家数据局", "智数科技", "合肥市数据资源局"]
    out_dir.mkdir(parents=True, exist_ok=True)
    (out_dir / "manifest.json").write_text(
        json.dumps(manifest, ensure_ascii=False, indent=2), encoding="utf-8")
    (out_dir / "whitelist.txt").write_text(
        "# 语料 v2 白名单示例\n" + "\n".join(manifest["whitelist"]) + "\n",
        encoding="utf-8")
    return manifest


def main(argv=None) -> int:
    ap = argparse.ArgumentParser("生成多类型合成测试语料 (v2)")
    ap.add_argument("--out", default="testdata_v2", help="输出目录")
    ap.add_argument("--per-type", type=int, default=12, help="每类文档数 (默认 12)")
    ap.add_argument("--seed", type=int, default=42, help="随机种子")
    args = ap.parse_args(argv)

    manifest = generate(Path(args.out), per_type=args.per_type, seed=args.seed)
    total_docs = len(manifest["docs"])
    total_truth = sum(len(d["entities"]) for d in manifest["docs"])
    print(f"[gen] 生成 {total_docs} 份文档 -> {args.out}")
    for dtype, n in manifest["types"].items():
        print(f"  {dtype:12} {n} 份")
    print(f"[gen] 真值实体 {total_truth} 条 (seed={args.seed})")
    return 0


if __name__ == "__main__":
    sys.exit(main())
